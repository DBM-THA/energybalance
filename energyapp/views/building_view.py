import csv
from io import BytesIO

from django.http import HttpResponse
from django.shortcuts import render, get_object_or_404, redirect
from energyapp.forms import BuildingForm,SimpleBuildingForm, Sheet01EnergyResultForm, EnergyResultSheet01Form
from energyapp.logic.building import calc_heating_demand
from energyapp.models import Building, Sheet01EnergyResult, EnergyResultSheet01
from energyapp.logic.result_sheet_01 import calculate_sheet01

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.barcharts import VerticalBarChart
from energyapp.forms import EnvelopeElementForm, LayerFormSet
from energyapp.models import EnvelopeElement, Building
from django.shortcuts import get_object_or_404
from energyapp.models import Building, EnvelopeElement, Layer






def dashboard(request):
    building_count = Building.objects.count()
    last_building = Building.objects.order_by("-id").first()
    return render(
        request,
        "energyapp/dashboard.html",
        {
            "building_count": building_count,
            "last_building": last_building,
        },
    )


def building_create_detailed(request):
    result = None

    if request.method == "POST":
        form = BuildingForm(request.POST)
        if form.is_valid():
            building = form.save()

            # Berechnung durchführen
            result = calc_heating_demand(building)
            building.ngf_t = result["ngf_t"]


            # Ergebnisse ins Modell schreiben
            building.result_Q_T = result["Q_T"]
            building.result_Q_V = result["Q_V"]
            building.result_Q_I = result["Q_I"]
            building.result_Q_S = result["Q_S"]
            building.result_Q_h = result["Q_h"]

            building.result_H_T = result["H_T"]
            building.result_H_V = result["H_V"]

            building.result_floor_area = result["floor_area"]
            building.result_roof_area = result["roof_area"]
            building.result_opaque_wall_area = result["opaque_wall_area"]
            building.result_window_area = result["window_area"]

            building.result_Q_S_n = result["Q_S_n"]
            building.result_Q_S_e = result["Q_S_e"]
            building.result_Q_S_s = result["Q_S_s"]
            building.result_Q_S_w = result["Q_S_w"]

            building.result_Q_PV_total = result["Q_PV_total"]
            building.result_Q_PV_on = result["Q_PV_on"]
            building.result_Q_PV_off = result["Q_PV_off"]

            building.save()

            # >>> HIER wird zur Gebäudeliste umgeleitet <<<
            return redirect("building_list")

    else:
        form = BuildingForm()

    context = {
        "form": form,
        "result": result,
        "edit_mode": False,
        "simple_mode": False,
    }
    return render(request, "energyapp/building_form.html", context)

def building_create_simple(request):
    if request.method == "POST":
        form = SimpleBuildingForm(request.POST)
        if form.is_valid():
            building = form.save()

            # Berechnung durchführen
            result = calc_heating_demand(building)
            building.ngf_t = result["ngf_t"]


            # Ergebnisse ins Modell schreiben
            building.result_Q_T = result["Q_T"]
            building.result_Q_V = result["Q_V"]
            building.result_Q_I = result["Q_I"]
            building.result_Q_S = result["Q_S"]
            building.result_Q_h = result["Q_h"]

            building.result_H_T = result["H_T"]
            building.result_H_V = result["H_V"]

            building.result_floor_area = result["floor_area"]
            building.result_roof_area = result["roof_area"]
            building.result_opaque_wall_area = result["opaque_wall_area"]
            building.result_window_area = result["window_area"]

            building.result_Q_S_n = result["Q_S_n"]
            building.result_Q_S_e = result["Q_S_e"]
            building.result_Q_S_s = result["Q_S_s"]
            building.result_Q_S_w = result["Q_S_w"]

            building.result_Q_PV_total = result["Q_PV_total"]
            building.result_Q_PV_on = result["Q_PV_on"]
            building.result_Q_PV_off = result["Q_PV_off"]

            building.save()
            return redirect("building_list")
    else:
        form = SimpleBuildingForm()

    return render(
        request,
        "energyapp/building_form.html",
        {
            "form": form,
            "edit_mode": False,
            "simple_mode": True,  # Flag für Template
        },
    )


def building_list(request):
    # erlaubte Sortierfelder
    order = request.GET.get("order", "-id")
    allowed_orders = {
        "name": "name",
        "name_desc": "-name",
        "area": "result_floor_area",
        "area_desc": "-result_floor_area",
        "q_h": "result_Q_h",
        "q_h_desc": "-result_Q_h",
        "id": "-id",
    }
    order_by = allowed_orders.get(order, "-id")

    buildings = Building.objects.all().order_by(order_by)
    context = {
        "buildings": buildings,
        "current_order": order,
    }
    return render(request, "energyapp/building_list.html", context)


def building_delete(request, pk):
    building = get_object_or_404(Building, pk=pk)
    if request.method == "POST":
        building.delete()
        return redirect("building_list")
    # falls jemand per GET draufgeht: einfach zurück zur Liste
    return redirect("building_list")


def building_delete_all(request):
    if request.method == "POST":
        Building.objects.all().delete()
        return redirect("building_list")
    return redirect("building_list")


def building_edit(request, pk):
    building = get_object_or_404(Building, pk=pk)

    if request.method == "POST":
        form = BuildingForm(request.POST, instance=building)
        if form.is_valid():
            building = form.save()

            # Berechnung erneut durchführen
            result = calc_heating_demand(building)
            building.ngf_t = result["ngf_t"]


            building.result_Q_T = result["Q_T"]
            building.result_Q_V = result["Q_V"]
            building.result_Q_I = result["Q_I"]
            building.result_Q_S = result["Q_S"]
            building.result_Q_h = result["Q_h"]

            building.result_H_T = result["H_T"]
            building.result_H_V = result["H_V"]

            building.result_floor_area = result["floor_area"]
            building.result_roof_area = result["roof_area"]
            building.result_opaque_wall_area = result["opaque_wall_area"]
            building.result_window_area = result["window_area"]

            building.result_Q_S_n = result["Q_S_n"]
            building.result_Q_S_e = result["Q_S_e"]
            building.result_Q_S_s = result["Q_S_s"]
            building.result_Q_S_w = result["Q_S_w"]

            building.result_Q_PV_total = result["Q_PV_total"]
            building.result_Q_PV_on = result["Q_PV_on"]
            building.result_Q_PV_off = result["Q_PV_off"]

            building.save()

            return redirect("building_detail", pk=building.pk)
    else:
        form = BuildingForm(instance=building)

    return render(
        request,
        "energyapp/building_form.html",
        {"form": form, "edit_mode": True, "building": building},
    )


def building_export_csv(request):
    """
    Export aller Gebäude als CSV.
    """
    buildings = Building.objects.all().order_by("id")

    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="buildings_export.csv"'

    writer = csv.writer(response, delimiter=";")

    # Kopfzeile
    writer.writerow([
        "ID",
        "Name",
        "Grundfläche [m²]",
        "Q_h [kWh/a]",
        "Q_T [kWh/a]",
        "Q_V [kWh/a]",
        "Q_I [kWh/a]",
        "Q_S [kWh/a]",
        "PV gesamt [kWh/a]",
        "PV Eigenverbrauch [kWh/a]",
        "PV Überschuss [kWh/a]",
    ])

    # Datenzeilen
    for b in buildings:
        writer.writerow([
            b.id,
            b.name,
            b.result_floor_area or "",
            b.result_Q_h or "",
            b.result_Q_T or "",
            b.result_Q_V or "",
            b.result_Q_I or "",
            b.result_Q_S or "",
            b.result_Q_PV_total or "",
            b.result_Q_PV_on or "",
            b.result_Q_PV_off or "",
        ])

    return response


def building_export_xlsx(request):
    """
    Export aller Gebäude als Excel-Datei (XLSX).
    """
    buildings = Building.objects.all().order_by("id")

    wb = Workbook()
    ws = wb.active
    ws.title = "Gebäude"

    # Kopfzeile
    headers = [
        "ID",
        "Name",
        "Grundfläche [m²]",
        "Q_h [kWh/a]",
        "Q_T [kWh/a]",
        "Q_V [kWh/a]",
        "Q_I [kWh/a]",
        "Q_S [kWh/a]",
        "PV gesamt [kWh/a]",
        "PV Eigenverb [kWh/a]",
        "PV Überschuss [kWh/a]",
    ]
    ws.append(headers)

    # Kopfzeile fett machen
    bold_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold_font

    # Datenzeilen
    for b in buildings:
        ws.append([
            b.id,
            b.name,
            b.result_floor_area or "",
            b.result_Q_h or "",
            b.result_Q_T or "",
            b.result_Q_V or "",
            b.result_Q_I or "",
            b.result_Q_S or "",
            b.result_Q_PV_total or "",
            b.result_Q_PV_on or "",
            b.result_Q_PV_off or "",
        ])

    # In Memory-Stream schreiben
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    # HTTP-Response
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="buildings_export.xlsx"'
    return response


def building_export_pdf(request):
    """
    Export aller Gebäude als übersichtliche PDF-Tabelle.
    """
    buildings = Building.objects.all().order_by("id")

    # PDF in Memory erzeugen
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=25 * mm,
        bottomMargin=20 * mm,
    )

    styles = getSampleStyleSheet()
    elements = []

    # Titel
    elements.append(Paragraph("Gebäude-Export – Energiebilanz", styles["Title"]))
    elements.append(Spacer(1, 8))

    # kleine Zusammenfassung
    info_text = f"Anzahl Gebäude: {buildings.count()}"
    elements.append(Paragraph(info_text, styles["Normal"]))
    elements.append(Spacer(1, 12))

    # Hilfsfunktion für Zahlenformat
    def fmt(val, decimals=1):
        if val is None:
            return "-"
        return f"{val:.{decimals}f}"

    # Tabellendaten
    data = [
        ["ID", "Name", "Grundfl. [m²]", "Q_h [kWh/a]", "PV on [kWh/a]", "PV off [kWh/a]"]
    ]

    for b in buildings:
        data.append([
            str(b.id),
            b.name,
            fmt(b.result_floor_area, 1),
            fmt(b.result_Q_h, 1),
            fmt(b.result_Q_PV_on, 1),
            fmt(b.result_Q_PV_off, 1),
        ])

    # Tabelle mit Styling
    table = Table(data, repeatRows=1, hAlign="LEFT")

    table.setStyle(TableStyle([
        # Header
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9ecef")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 10),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
        ("TOPPADDING", (0, 0), (-1, 0), 6),

        # Körper
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 4),
        ("TOPPADDING", (0, 1), (-1, -1), 4),

        # Gitterlinien
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),

        # Ausrichtung: ID + Zahlen rechts, Name links
        ("ALIGN", (0, 1), (0, -1), "RIGHT"),
        ("ALIGN", (2, 1), (-1, -1), "RIGHT"),

        # wechselnde Zeilenfarben im Body
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#f8f9fa")]),
    ]))

    elements.append(table)

    # PDF aufbauen
    doc.build(elements)

    # Response zurückgeben
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/pdf"
    )
    response["Content-Disposition"] = 'attachment; filename=\"buildings_export.pdf\"'
    return response


def building_result_pdf(request, pk):
    """
    PDF-Bericht für EIN Gebäude inkl. Tabelle und einfachem Balkendiagramm.
    """
    building = get_object_or_404(Building, pk=pk)

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=25 * mm,
        bottomMargin=20 * mm,
    )

    styles = getSampleStyleSheet()
    elements = []

    # Titel
    title = f"Energiebericht – {building.name}"
    elements.append(Paragraph(title, styles["Title"]))
    elements.append(Spacer(1, 10))

    # Kurzinfo
    info_lines = [
        f"Gebäude-ID: {building.id}",
        f"Grundfläche: {building.result_floor_area or '-'} m²",
        f"Raum-Solltemperatur: {building.setpoint_temp or '-'} °C",
    ]
    for line in info_lines:
        elements.append(Paragraph(line, styles["Normal"]))
    elements.append(Spacer(1, 12))

    # Hilfsfunktion Zahlenformat
    def fmt(val, decimals=1):
        if val is None:
            return "-"
        return f"{val:.{decimals}f}"

    # Ergebnis-Tabelle
    data = [
        ["Größe", "Wert"],
        ["Heizwärmebedarf Q_h [kWh/a]", fmt(building.result_Q_h, 1)],
        ["Transmissionsverluste Q_T [kWh/a]", fmt(building.result_Q_T, 1)],
        ["Lüftungsverluste Q_V [kWh/a]", fmt(building.result_Q_V, 1)],
        ["Interne Gewinne Q_I [kWh/a]", fmt(building.result_Q_I, 1)],
        ["Solare Gewinne Q_S [kWh/a]", fmt(building.result_Q_S, 1)],
        ["PV-Gesamt [kWh/a]", fmt(building.result_Q_PV_total, 1)],
        ["PV-Eigenverbrauch [kWh/a]", fmt(building.result_Q_PV_on, 1)],
        ["PV-Überschuss [kWh/a]", fmt(building.result_Q_PV_off, 1)],
    ]

    table = Table(data, hAlign="LEFT", colWidths=[70 * mm, 60 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9ecef")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(Paragraph("Ergebnisse – Übersicht", styles["Heading2"]))
    elements.append(table)
    elements.append(Spacer(1, 16))

    # Balkendiagramm Energieflüsse
    elements.append(Paragraph("Energieflüsse (jährlich)", styles["Heading2"]))
    elements.append(Spacer(1, 4))

    flows = [
        ("Q_T", building.result_Q_T or 0),
        ("Q_V", building.result_Q_V or 0),
        ("Q_I", building.result_Q_I or 0),
        ("Q_S", building.result_Q_S or 0),
        ("Q_h", building.result_Q_h or 0),
    ]

    data_values = [[v for (_, v) in flows]]
    labels = [lbl for (lbl, _) in flows]

    drawing = Drawing(400, 200)
    bc = VerticalBarChart()
    bc.x = 40
    bc.y = 40
    bc.height = 120
    bc.width = 320
    bc.data = data_values
    bc.categoryAxis.categoryNames = labels
    bc.categoryAxis.labels.boxAnchor = "n"
    bc.valueAxis.valueMin = 0
    bc.barWidth = 20
    bc.groupSpacing = 10
    bc.barSpacing = 5
    bc.bars[0].fillColor = colors.HexColor("#0d6efd")

    drawing.add(bc)
    elements.append(drawing)

    # PDF bauen
    doc.build(elements)

    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/pdf",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="building_{building.id}_bericht.pdf"'
    )
    return response

def internal_gains(request):
    # Platzhalter – später durch richtigen Inhalt ersetzen (Derya, Lucy)
    return render(request, "energyapp/internal_gains.html")


def envelope_detail(request):
    COMPONENT_CHOICES = {
        "kellerdecke": "Kellerdecke",
        "dach": "Dach",
        "aw_nord": "Massive Außenwand Nord",
        "aw_sued": "Massive Außenwand Süd",
        "aw_ost": "Massive Außenwand Ost",
        "aw_west": "Massive Außenwand West",
        "openings": "Fenster / Tür",
    }

    # Auswahl
    component = request.GET.get("component") or request.POST.get("component") or ""
    component_label = COMPONENT_CHOICES.get(component, "")

    # building_id optional (wenn ihr’s habt: perfekt)
    building_id = request.GET.get("building_id") or request.POST.get("building_id") or ""
    building = Building.objects.filter(pk=building_id).first() if building_id else None

    # Rsi/Rse abhängig vom Bauteil
    RSI_BY_FLOW = {"up": 0.10, "horizontal": 0.13, "down": 0.17}
    RSE = 0.04
    COMPONENT_TO_FLOW = {
        "dach": "up",
        "kellerdecke": "down",
        "aw_nord": "horizontal",
        "aw_sued": "horizontal",
        "aw_ost": "horizontal",
        "aw_west": "horizontal",
        "openings": "horizontal",
    }
    heat_flow = COMPONENT_TO_FLOW.get(component, "horizontal")
    R_FIXED = float(RSI_BY_FLOW[heat_flow] + RSE)

    # Bearbeiten nur, wenn element_id explizit kommt
    element_id = request.GET.get("element_id") or request.POST.get("element_id")
    element = None
    if element_id:
        if building:
            element = get_object_or_404(EnvelopeElement, pk=element_id, building=building)
        else:
            element = get_object_or_404(EnvelopeElement, pk=element_id)

    # -------- POST: speichern --------
    if request.method == "POST":
        # Ohne Auswahl nicht speichern
        if not component_label:
            return redirect(request.path)

        post = request.POST.copy()

        # Name automatisch setzen
        if not post.get("name"):
            post["name"] = component_label

        # ✅ layer_type serverseitig setzen (weil im Template versteckt)
        prefix = LayerFormSet.get_default_prefix()  # meistens "layer_set" oder "layers" je nach Django
        total = int(post.get(f"{prefix}-TOTAL_FORMS", "0") or "0")
        for i in range(total):
            post[f"{prefix}-{i}-layer_type"] = "layer"

        form = EnvelopeElementForm(post, instance=element)
        formset = LayerFormSet(post, instance=element)

        if form.is_valid() and formset.is_valid():
            element = form.save(commit=False)

            # ✅ wenn building genutzt wird: korrekt verknüpfen
            if building:
                element.building = building

            element.heat_flow = heat_flow

            # ✅ Name fix auf Dropdown-Label (damit es konsistent ist)
            element.name = component_label

            element.save()

            # Formset speichern (aber leere Zeilen ignorieren)
            objs = formset.save(commit=False)
            saved_any = False
            for obj in objs:
                obj.element = element
                if (
                    (obj.name and obj.name.strip())
                    or (obj.thickness is not None)
                    or (obj.lambda_value is not None)
                    or (obj.R_value is not None)
                ):
                    obj.save()
                    saved_any = True

            for obj in formset.deleted_objects:
                obj.delete()

            # U-Wert berechnen (mit richtigem R_FIXED)
            if element.use_custom_layers and saved_any:
                element.calculate_u_value(R_fixed=R_FIXED)
            elif (not element.use_custom_layers) and (element.u_value_external is not None):
                element.u_value_calculated = element.u_value_external
                element.save(update_fields=["u_value_calculated"])
            else:
                element.u_value_calculated = None
                element.save(update_fields=["u_value_calculated"])

            # ✅ PRG Redirect => danach ist Formular nicht mehr „gefüllt“
            url = request.path + "?component=" + component + "&just_saved=1"
            if building_id:
                url += "&building_id=" + str(building_id)
            return redirect(url)

        # Wenn ungültig: Seite mit Fehlern zurückgeben
        saved_elements = _get_saved_elements(building)
        return render(request, "energyapp/envelope.html", {
            "component_choices": COMPONENT_CHOICES,
            "selected_component": component,
            "selected_building": building,
            "building_id": building_id,

            "selected_element": element,
            "form": form,
            "formset": formset,

            "R_FIXED": R_FIXED,
            "saved_u": (element.u_value_calculated if element else None),

            "show_saved_list": True,
            "saved_elements": saved_elements,

            "ui_error": "Speichern fehlgeschlagen – siehe rote Box (Form Errors).",
        })

    # -------- GET: anzeigen --------
    just_saved = (request.GET.get("just_saved") == "1")

    # Formular: nach Save wieder „frisch“ (leeres Formset), außer wenn Bearbeiten
    if element_id and element:
        form = EnvelopeElementForm(instance=element)
        formset = LayerFormSet(instance=element)
        saved_u = element.u_value_calculated
    else:
        initial = {"name": component_label} if component_label else {}
        form = EnvelopeElementForm(initial=initial)
        formset = LayerFormSet()
        saved_u = None

    saved_elements = _get_saved_elements(building) if just_saved else []

    return render(request, "energyapp/envelope.html", {
        "component_choices": COMPONENT_CHOICES,
        "selected_component": component,
        "selected_building": building,
        "building_id": building_id,

        "selected_element": element if element_id else None,
        "form": form,
        "formset": formset,

        "R_FIXED": R_FIXED,
        "saved_u": saved_u,

        # ✅ Liste nur nach erfolgreichem Save
        "show_saved_list": just_saved,
        "saved_elements": saved_elements,
        "ui_error": "",
    })


def _get_saved_elements(building):
    qs = EnvelopeElement.objects.all()

    # ✅ wichtig: sonst kommen „Kellerdecke“/alte Tests/andere Gebäude rein
    if building:
        qs = qs.filter(building=building)

    qs = qs.order_by("-id").prefetch_related("layers")

    # ✅ nur Elemente, die wirklich Inhalt haben
    ids_with_layers = set(Layer.objects.values_list("element_id", flat=True).distinct())
    return [
        el for el in qs
        if (el.id in ids_with_layers) or (el.u_value_external is not None) or (el.u_value_calculated is not None)
    ]








def solar_gains(request):
    # Platzhalter – später durch richtigen Inhalt ersetzen
    return render(request, "energyapp/solar_gains.html")

def ventilation(request):
    # Platzhalter – später durch richtigen Inhalt ersetzen
    return render(request, "energyapp/ventilation.html")

def building_detail(request, pk):
    """
    Detailansicht eines Gebäudes mit gespeichertem Ergebnis.
    (Template verwenden wir später, wenn du soweit bist.)
    """
    building = get_object_or_404(Building, pk=pk)
    return render(request, "energyapp/building_detail.html", {"building": building})

def summary_dashboard(request):
    building_id = request.GET.get("building")
    if building_id:
        building = get_object_or_404(Building, pk=building_id)
    else:
        building = Building.objects.order_by("-id").first()


    if not building:
        return render(
            request,
            "energyapp/summary_dashboard.html",
            {
                "building": None,
                "sheet": None,
                "calc": {},
                "form": None,
            },
        )

    sheet, _ = EnergyResultSheet01.objects.get_or_create(building=building)

    if request.method == "POST":
        form = EnergyResultSheet01Form(request.POST, instance=sheet)
        if form.is_valid():
            sheet = form.save()
            return redirect("summary_dashboard")
    else:
        form = EnergyResultSheet01Form(instance=sheet)

    calc = calculate_sheet01(building, sheet)

    return render(
        request,
        "energyapp/summary_dashboard.html",
        {
            "building": building,
            "sheet": sheet,
            "calc": calc,
            "form": form,
        },
    )

def pv_details(request):
    return render(request, "energyapp/pv_details.html")
